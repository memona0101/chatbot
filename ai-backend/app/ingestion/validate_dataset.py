import openpyxl

DATASET = "data/MoinSystems_AI_Public_Chatbot_RAG_Dataset_v2 (1).xlsx"
SHEET = "RAG_Knowledge"

REQUIRED_FIELDS = [
    "ID",
    "Title",
    "Category",
    "Tags",
    "Intents",
    "Content",
]


def validate_dataset():
    workbook = openpyxl.load_workbook(
        DATASET,
        read_only=True,
        data_only=True,
    )

    if SHEET not in workbook.sheetnames:
        raise ValueError(f"Missing sheet: {SHEET}")

    sheet = workbook[SHEET]

    rows = sheet.iter_rows(values_only=True)
    first_row = list(next(rows))

    cleaned_headers = [
        str(header).strip()
        for header in first_row
        if header is not None and str(header).strip()
    ]

    print("Headers:", cleaned_headers)

    missing_fields = [
        field
        for field in REQUIRED_FIELDS
        if field not in cleaned_headers
    ]

    if missing_fields:
        raise ValueError(
            f"Missing required fields: {missing_fields}"
        )

    header_indexes = {
        str(header).strip(): index
        for index, header in enumerate(first_row)
        if header is not None and str(header).strip()
    }

    record_count = 0
    ids = []
    errors = []

    for row_number, row in enumerate(rows, start=2):

        # Ignore completely empty Excel rows
        if not any(
            value is not None and str(value).strip()
            for value in row
        ):
            continue

        record_count += 1

        record_id = row[header_indexes["ID"]]

        if not record_id:
            errors.append(f"Row {row_number}: missing ID")
        else:
            ids.append(str(record_id).strip())

        for field in REQUIRED_FIELDS:
            value = row[header_indexes[field]]

            if value is None or not str(value).strip():
                errors.append(
                    f"Row {row_number}: missing {field}"
                )

    duplicate_ids = {
        record_id
        for record_id in ids
        if ids.count(record_id) > 1
    }

    print(f"Total records: {record_count}")
    print(f"Unique IDs: {len(set(ids))}")
    print(f"Duplicate IDs: {len(duplicate_ids)}")
    print(f"Validation errors: {len(errors)}")

    if duplicate_ids:
        print("Duplicate IDs:")
        for duplicate_id in sorted(duplicate_ids):
            print(f"  - {duplicate_id}")

    if errors:
        print("\nValidation errors:")
        for error in errors[:20]:
            print(f"  - {error}")

    if errors or duplicate_ids:
        raise SystemExit("Dataset validation FAILED")

    print("\nDataset validation PASSED")


if __name__ == "__main__":
    validate_dataset()
    